import re

import pandas as pd
import numpy as np
from pathlib import Path
from openpyxl import load_workbook

INPUT_PATH = "Data/raw/Tablet masses(1) 2.xlsx"
OUTPUT_RAW_CSV = "Data/processed/raw_masses.csv"
OUTPUT_SUMMARY_CSV = "Data/processed/summary_stats.csv"

N_TABLETS = 20

GEOMETRIES = [
    "Rectangular cranked 10mm",
    "Rectangular cranked 5.5mm",
    "Round cranked 10mm",
    "Round cranked 5.5mm",
    "Spider geometry",
    "Cut-out geometry",
]

POWDER_SHEET_MAP = {
    "DCP": "DCP",
    "TCC": "TCC",
    "Lactose Anhydrate": "Lactose Anhydrate",
    "MCC coarse": "MCC coarse",
    "MCC fine": "MCC fine",
    "Powdered cellulose(ArbocelP290)": "Powdered cellulose(ArbocelP290)",
    "MCC finest (Vivapur102)": "MCC finest (Vivapur102)",
    #"MCC finest +Pellets ": "MCC finest +Pellets ",
    "HPMC(60%Tylopur40%V120)": "HPMC(60%Tylopur40%V120)",
    "CaCO3(50%MagGranCC)": "CaCO3(50%MagGranCC)",
}

# Unicode subscript translation map
SUB_MAP = str.maketrans("ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ", "aehijklmnoprstuvx")

EXPECTED_NT_NP = {
    25: {13, 25, 38},
    63: {31, 63, 93},
    100: {25, 38, 50, 75, 100, 150},
}
EXPECTED_ROWS_PER_BLOCK = sum(len(v) for v in EXPECTED_NT_NP.values())

EXPECTED_TABLETS = 20

MEAN_ATOL = 0.01
STD_ATOL = 0.01
RSD_ATOL = 0.001



KEYS = ["powder", "geometry", "nt", "np"]

def _norm(s):
    """lowercase, strip, collapse whitespace/punctuation, and translate subscripts"""
    if s is None:
        return ""
    s_str = str(s).strip().lower().translate(SUB_MAP)
    return re.sub(r"[\s\.\-\[\]\(\)]+", "", s_str)


GEOMETRY_LOOKUP = {_norm(g): g for g in GEOMETRIES}
assert len(GEOMETRY_LOOKUP) == len(GEOMETRIES), "Geometry name collide afetr normalising"

def _to_num(v):
    """Return v as int/float, or None if it isn't numeric."""
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    if np.isnan(x):
        return None
    return int(x) if x.is_integer() else x
 
 
def _issue(powder, geometry, problem, detail=""):
    return {"powder": powder, "geometry": geometry, "problem": problem, "detail": detail}
 
# Locating Blocks
def find_geometry_blocks(ws):
    """Return list of (geometry_name, title_row_idx) in row order."""
    blocks = []
    for row in ws.iter_rows():
        for cell in row:
            key = _norm(cell.value)
            if key in GEOMETRY_LOOKUP:
                blocks.append((GEOMETRY_LOOKUP[key], cell.row))
                break
    return blocks


def find_header_row(ws, start_row, max_lookahead=10):
    for r in range(start_row + 1, start_row + 1 + max_lookahead):
        row_vals = [c.value for c in ws[r] if c.value is not None]
        joined = " ".join(_norm(v) for v in row_vals)
        if ("tab" in joined or "1" in joined) and ("nt" in joined or "n1" in joined or "n" in joined):
            return r
    return None


def map_columns(ws, header_row):
    """Return dict: {'nt': col_idx, 'np': col_idx, 'tabs': {1: col, ...},
    'mean': col_idx, 'std': col_idx, 'rel': col_idx}"""
    cols = {"tabs": {}}
    for cell in ws[header_row]:
        if cell.value is None:
            continue
        key = _norm(cell.value)
        tab_match = re.match(r"tab(\d+)", key)
        if tab_match:
            cols["tabs"][int(tab_match.group(1))] = cell.column
        elif key in ("nt", "n1", "ntrpm", "n1rpm"):
            cols["nt"] = cell.column
        elif key in ("np", "nprpm"):
            cols["np"] = cell.column
        elif "rel" in key:
            cols["rel"] = cell.column
        elif "std" in key:
            cols["std"] = cell.column
        elif "mean" in key:
            cols["mean"] = cell.column
    return cols


def row_is_blank(ws, r, cols_to_check):
    return all(ws.cell(row=r, column=c).value in (None, "") for c in cols_to_check)

# Parsing
def parse_block(ws, powder, geometry, header_row, next_boundary_row):
    cols = map_columns(ws, header_row)
    issues = []
    if "nt" not in cols or "np" not in cols or not cols["tabs"]:
        issues.append(_issue(powder, geometry, "missing nt/np/Tab columns", f"header row {header_row}"))
        return [], [], issues

    if len(cols["tabs"]) != EXPECTED_TABLETS:
        issues.append(
            _issue(powder, geometry, "unexpected number of Tab columns", f"found {len(cols['tabs'])}")
        )
 
    expected = {(nt, np_) for nt, nps in EXPECTED_NT_NP.items() for np_ in nps}
    seen = set()
    raw_rows, summary_rows = [], []
    last_nt = None
 
    for r in range(header_row + 1, next_boundary_row):
        nt_num = _to_num(ws.cell(row=r, column=cols["nt"]).value)
        np_cell = ws.cell(row=r, column=cols["np"]).value
 
        # nt is written once per group, so forward-fill it. Only update on a
        # genuine number so a text note can't wipe it out.
        if nt_num is not None:
            last_nt = nt_num
 
        if np_cell in (None, ""):
            continue  # blank row or nt-only row
 
        key = (last_nt, _to_num(np_cell))
        if key not in expected:
            issues.append(
                _issue(powder, geometry, "unexpected (nt,np) row skipped", f"row {r}: nt={last_nt}, np={np_cell}")
            )
            continue
        if key in seen:
            issues.append(
                _issue(powder, geometry, "duplicate (nt,np) row skipped", f"row {r}: nt={key[0]}, np={key[1]}")
            )
            continue
        seen.add(key)
 
        # Keep a row for EVERY tablet column, even if empty, so gaps stay
        # visible as NaN instead of disappearing.
        for tab_n, col in sorted(cols["tabs"].items()):
            mass = ws.cell(row=r, column=col).value
            if mass == "":
                mass = None
            raw_rows.append(
                {
                    "powder": powder,
                    "geometry": geometry,
                    "nt": key[0],
                    "np": key[1],
                    "tablet_id": tab_n,
                    "mass_mg": mass,
                }
            )
 
        summary_rows.append(
            {
                "powder": powder,
                "geometry": geometry,
                "nt": key[0],
                "np": key[1],
                "mean_mg": ws.cell(row=r, column=cols["mean"]).value if "mean" in cols else None,
                "std_mg": ws.cell(row=r, column=cols["std"]).value if "std" in cols else None,
                "rel_std_pct": ws.cell(row=r, column=cols["rel"]).value if "rel" in cols else None,
            }
        )
 
    for nt, np_ in sorted(expected - seen):
        issues.append(_issue(powder, geometry, "expected (nt,np) row missing", f"nt={nt}, np={np_}"))
 
    return raw_rows, summary_rows, issues


def parse_workbook(path):
    wb = load_workbook(path, data_only=True)
    all_raw, all_summary, all_issues = [], [], []
 
    for ws in wb.worksheets:
        if ws.title not in POWDER_SHEET_MAP:
            print(f"[sheet] '{ws.title}': not in POWDER_SHEET_MAP, skipping")
            continue
        powder = POWDER_SHEET_MAP[ws.title]
 
        blocks = find_geometry_blocks(ws)
        found = {g for g, _ in blocks}
        for g in GEOMETRIES:
            if g not in found:
                all_issues.append(_issue(powder, g, "geometry title not found in sheet"))
        if not blocks:
            print(f"[sheet] {ws.title}: no geometry blocks found, skipping")
            continue
        print(f"[sheet] {ws.title} -> powder='{powder}': {len(blocks)} blocks found")
 
        for i, (geometry, title_row) in enumerate(blocks):
            header_row = find_header_row(ws, title_row)
            if header_row is None:
                all_issues.append(_issue(powder, geometry, "no header row found", f"below row {title_row}"))
                continue
            next_boundary = blocks[i + 1][1] if i + 1 < len(blocks) else ws.max_row + 1
            raw_rows, summary_rows, issues = parse_block(ws, powder, geometry, header_row, next_boundary)
            print(f"  [{'ok' if not issues else 'CHECK'}] {geometry}: {len(summary_rows)} (nt,np) rows")
            all_raw.extend(raw_rows)
            all_summary.extend(summary_rows)
            all_issues.extend(issues)
 
    return pd.DataFrame(all_raw), pd.DataFrame(all_summary), pd.DataFrame(all_issues)
 
 
def coerce_numeric(df, cols, id_cols):
    """Convert `cols` to numeric; print + return any cells that failed."""
    problems = []
    for col in cols:
        if col not in df.columns:
            continue
        raw = df[col]
        converted = pd.to_numeric(raw, errors="coerce")
        bad_mask = converted.isna() & raw.notna() & (raw != "")
        if bad_mask.any():
            bad = df.loc[bad_mask, id_cols + [col]].copy()
            bad["column"] = col
            bad = bad.rename(columns={col: "bad_value"})
            problems.append(bad)
        df[col] = converted
    if problems:
        report = pd.concat(problems, ignore_index=True)
        print(f"\nWarning: {len(report)} non-numeric value(s) found and set to NaN:")
        print(report.to_string(index=False))
        return df, report
    return df, pd.DataFrame()
 
# ------------------------------------------------------------------ VALIDATION
 
 
def add_n_tablets(raw_df, summary_df):
    """Number of actual (non-missing) tablet masses behind each condition."""
    counts = raw_df.groupby(KEYS)["mass_mg"].count().rename("n_tablets").reset_index()
    return summary_df.merge(counts, on=KEYS, how="left")

def check_counts(raw_df, summary_df):
    n_powders = len(set(POWDER_SHEET_MAP.values()))
    expected_rows = n_powders * len(GEOMETRIES) * EXPECTED_ROWS_PER_BLOCK
    status = "OK" if len(summary_df) == expected_rows else "MISMATCH"
    print(f"\n[{status}] summary rows: {len(summary_df)} (expected {expected_rows} = "
          f"{n_powders} powders x {len(GEOMETRIES)} geometries x {EXPECTED_ROWS_PER_BLOCK})")
 
    per_block = summary_df.groupby(["powder", "geometry"]).size()
    bad_blocks = per_block[per_block != EXPECTED_ROWS_PER_BLOCK]
    if len(bad_blocks):
        print(f"[MISMATCH] blocks without {EXPECTED_ROWS_PER_BLOCK} rows:")
        print(bad_blocks.to_string())
 
    short = summary_df[summary_df["n_tablets"] != EXPECTED_TABLETS]
    if len(short):
        print(f"[INFO] {len(short)} condition(s) with n_tablets != {EXPECTED_TABLETS} (missing masses kept as NaN):")
        print(short[KEYS + ["n_tablets"]].to_string(index=False))
    else:
        print(f"[OK] every condition has {EXPECTED_TABLETS} tablets")
 
 
def check_summary_consistency(raw_df, summary_df):
    """Compare the sheet's hand-calculated Mean/Std/RSD with values recomputed
    from the tablet masses. FLAGS differences; never overwrites anything."""
    stats = (
        raw_df.groupby(KEYS)["mass_mg"]
        .agg(
            mean_recalc="mean",
            std_recalc_ddof1=lambda s: s.std(ddof=1),
            std_recalc_ddof0=lambda s: s.std(ddof=0),
        )
        .reset_index()
    )
    m = summary_df.merge(stats, on=KEYS, how="left")
 
    m["mean_ok"] = np.isclose(m["mean_mg"], m["mean_recalc"], atol=MEAN_ATOL, rtol=0)
    std1 = np.isclose(m["std_mg"], m["std_recalc_ddof1"], atol=STD_ATOL, rtol=0)
    std0 = np.isclose(m["std_mg"], m["std_recalc_ddof0"], atol=STD_ATOL, rtol=0)
    m["std_ok"] = std1 | std0

    print("matches ddof=1:", std1.sum())
    print("matches ddof=0:", std0.sum())
    print("matches both: ", (std1 & std0).sum())
    print("ddof=0 only: ", (std0 & ~std1).sum())

    m["std_formula"] = np.where(std1, "ddof=1", np.where(std0, "ddof=0", "neither"))
    ddof0 = m[m["std_formula"] == "ddof=0"]
    print("\nRows where the sheet used ddof=0:")
    print(ddof0.groupby(["powder", "geometry"]).size())
    ddof0[KEYS + ["std_mg", "std_recalc_ddof1", "std_recalc_ddof0"]].to_csv(
    "Data/processed/std_ddof0_rows.csv", index=False
    )
    # RSD uses only the sheet's own columns: Std / Mean * 100
    m["rsd_recalc"] = m["std_mg"] / m["mean_mg"] * 100
    m["rsd_ok"] = np.isclose(m["rsd_pct"], m["rsd_recalc"], atol=RSD_ATOL, rtol=0)
 
    print(f"\nStd formula used in sheet: matches ddof=1 in {std1.mean():.1%} of rows, "
          f"ddof=0 in {std0.mean():.1%}")
    flagged = m[~(m["mean_ok"] & m["std_ok"] & m["rsd_ok"])]
    if flagged.empty:
        print("[OK] sheet Mean/Std/RSD are consistent with the tablet masses")
    else:
        print(f"[CHECK] {len(flagged)} condition(s) where sheet values disagree with the masses "
              f"(or contain NaN) - see summary_consistency_flags.csv")
    return flagged
 
 
# ------------------------------------------------------------------ RUN
 
if __name__ == "__main__":
    raw_df, summary_df, issues_df = parse_workbook(INPUT_PATH)
 
    print(f"\nraw_masses: {len(raw_df)} rows")
    print(f"summary_stats: {len(summary_df)} rows")
 
    # Clean up mixed text/number cells before doing any math on them.
    raw_df, raw_bad = coerce_numeric(raw_df, ["mass_mg"], KEYS + ["tablet_id"])
    summary_df, summary_bad = coerce_numeric(summary_df, ["mean_mg", "std_mg", "rel_std_pct"], KEYS)

    assert summary_df["rel_std_pct"].max() < 1, "mixed units? some RSD values look like percentages"
    summary_df["rsd_pct"] = summary_df["rel_std_pct"] * 100
    summary_df = summary_df.drop(columns="rel_std_pct")

    for p in (OUTPUT_RAW_CSV, OUTPUT_SUMMARY_CSV):
            Path(p).parent.mkdir(parents=True, exist_ok=True)  # creates Data/processed if missing
    
    if not raw_bad.empty:
        raw_bad.to_csv("Data/processed/raw_masses_bad_cells.csv", index=False)
    if not summary_bad.empty:
        summary_bad.to_csv("Data/processed/summary_stats_bad_cells.csv", index=False)
 
    summary_df = add_n_tablets(raw_df, summary_df)
 
    # Structural problems found while parsing (stray rows, missing rows, ...)
    if issues_df.empty:
        print("\n[OK] no parsing issues")
    else:
        print(f"\n[CHECK] {len(issues_df)} parsing issue(s):")
        print(issues_df.to_string(index=False))
        issues_df.to_csv("Data/processed/parse_issues.csv", index=False)
 
    check_counts(raw_df, summary_df)
    flagged = check_summary_consistency(raw_df, summary_df)
    if not flagged.empty:
        flagged.to_csv("Data/processed/summary_consistency_flags.csv", index=False)
        print(flagged[["mean_ok", "std_ok", "rsd_ok"]].mean())
        print(flagged[["mean_mg", "mean_recalc", "std_mg", "std_recalc_ddof1", "rsd_pct", "rsd_recalc"]].head(10))
 
    raw_df.to_csv(OUTPUT_RAW_CSV, index=False)
    summary_df.to_csv(OUTPUT_SUMMARY_CSV, index=False)
    print(f"\nWrote {OUTPUT_RAW_CSV}, {OUTPUT_SUMMARY_CSV}")
 